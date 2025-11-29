from docx import Document
import pandas as pd
from datetime import datetime

# Rutas de archivo (ajusta según tu carpeta)
PATH_CSV = "clientes.csv"
PATH_TEMPLATE = "plantilla.docx"
OUTPUT_FOLDER = "cartas"  # opcional: carpeta de salida

import os
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# Leer CSV (separador ;)
df = pd.read_csv(PATH_CSV, sep=";")

# Fecha actual en formato europeo
fecha_hoy = datetime.today().strftime("%d/%m/%Y")

for _, row in df.iterrows():
    doc = Document(PATH_TEMPLATE)

    # Reemplazo en párrafos
    for p in doc.paragraphs:
        if "{{FECHA}}" in p.text:
            p.text = p.text.replace("{{FECHA}}", fecha_hoy)
        if "{{NOMBRE}}" in p.text:
            p.text = p.text.replace("{{NOMBRE}}", str(row["nombre"]))
        if "{{DIRECCION}}" in p.text:
            p.text = p.text.replace("{{DIRECCION}}", str(row["direccion"]))
        if "{{CIUDAD}}" in p.text:
            p.text = p.text.replace("{{CIUDAD}}", str(row["ciudad"]))
        if "{{PAIS}}" in p.text:
            p.text = p.text.replace("{{PAIS}}", str(row["pais"]))

    # (Opcional) si tienes tablas con placeholders, también se podría iterar aquí

    # Guardar carta personalizada
    output_path = os.path.join(OUTPUT_FOLDER, f"carta_{row['id']}.docx")
    doc.save(output_path)

print("Cartas generadas correctamente.")
